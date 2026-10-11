"""Check both report renderers and the declared Song/1.5 typography."""
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import statistics

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
import pdfplumber

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliveries/connected_topk0175_20261003"
DOCX = OUT / "量价分工集成股票收益预测分析报告.docx"
PDF = OUT / "量价分工集成股票收益预测分析报告.pdf"


def main():
    doc = Document(DOCX)
    started = False
    checked = 0
    for p in doc.paragraphs:
        if p.text == "摘要":
            started = True
        if not started or not p.text:
            continue
        assert p.paragraph_format.line_spacing == 1.5, p.text[:30]
        for run in p.runs:
            if run.text:
                assert run.font.name == "宋体" and run.font.size.pt == 10.5, p.text[:30]
                assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
        checked += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    assert p.paragraph_format.line_spacing == 1.5
                    for run in p.runs:
                        if run.text:
                            assert run.font.name == "宋体" and run.font.size.pt == 10.5
    renderers = {}
    for name, path in [("Word", OUT / "qa/word_report.pdf"), ("LibreOffice", PDF)]:
        reader = PdfReader(path)
        texts = [p.extract_text() or "" for p in reader.pages]
        assert len(texts) == 10
        assert "摘要" in texts[1] and "参考文献" in texts[9]
        assert all(len(t.strip()) > 40 for t in texts)
        assert "0.445916" in texts[1] and "0.414869" in texts[7]
        assert "0.417099" in texts[7]
        with pdfplumber.open(path) as pdf:
            page = pdf.pages[5]  # a prose-only body page
            chars = [c for c in page.chars if abs(c["size"]-10.5) < 0.05 and 40 < c["top"] < 760]
            tops = Counter(round(c["top"], 1) for c in chars)
            lines = sorted(t for t, count in tops.items() if count >= 15)
            deltas = [b-a for a, b in zip(lines, lines[1:]) if 20 < b-a < 25]
            spacing = statistics.median(deltas)
            if name == "LibreOffice":
                assert abs(spacing - 22.2) <= 0.15, (name, spacing)
            else:
                assert 20 <= spacing <= 23, (name, spacing)
            assert len(deltas) >= 10
        renderers[name] = {"pages": 10, "body_pages": 8, "measured_line_spacing_pt": spacing,
                           "measured_gaps": len(deltas), "pdf_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    with ZipFile(DOCX) as package:
        assert b"w:pBdr" not in package.read("word/document.xml")
    result = {"status": "structural_checks_passed", "body_font": "宋体10.5pt",
              "line_spacing": 1.5, "body_paragraphs_checked": checked, "tables": len(doc.tables),
              "figures": len(doc.inline_shapes), "renderers": renderers,
              "docx_sha256": hashlib.sha256(DOCX.read_bytes()).hexdigest(),
              "pdf_renderer": "LibreOffice; delivered PDF measures22.2pt. Word uses the same1.5 OOXML spacing but measures20.4pt with its font engine; both have8 body pages.",
              "visual_review": "pending final all-page inspection"}
    (OUT / "report_audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
