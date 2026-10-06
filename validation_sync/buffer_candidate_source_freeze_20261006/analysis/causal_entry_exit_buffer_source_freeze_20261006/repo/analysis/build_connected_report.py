"""Write the final report using the retained academic report's page system.

Use the Codex bundled Python; render with the packaged render_docx.py afterwards.
"""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
ARCHIVE = ROOT.parent / "赛题五_历史归档_20261001"
HERE = ROOT / "analysis/connected_validation_20261002"
OUT = ROOT / "deliveries/connected_topk0175_20261003"
REF = ARCHIVE / "deliveries/peer67_slow16_重写版_20260930/XGBoost67特征16种子数据分析报告.docx"
TARGET = OUT / "量价分工集成股票收益预测分析报告.docx"
QA = OUT / "qa"


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def charts():
    QA.mkdir(exist_ok=True)
    plt.rcParams.update({"font.family": "Microsoft YaHei", "font.size": 9,
                         "axes.spines.top": False, "axes.spines.right": False,
                         "axes.unicode_minus": False, "savefig.dpi": 200})
    curve = read(HERE / "wave36_results.json")["curve_ordered"]
    fig, ax = plt.subplots(figsize=(6.5, 1.9))
    xx = [r["frac"] * 100 for r in curve]
    yy = [r["connected_final_score"] for r in curve]
    ax.plot(xx, yy, "o-", color="#26445d", linewidth=1.5)
    ax.scatter([17.5], [max(yy)], color="#ac4b2b", s=48, zorder=3)
    ax.annotate(f"17.5%  {max(yy):.6f}", (17.5, max(yy)), xytext=(4, -20), textcoords="offset points")
    ax.set(xlabel="训练损失中的 Top-K 比例（%）", ylabel="连接验证综合分", ylim=(0.4418, 0.4464))
    ax.grid(axis="y", alpha=0.2)
    fig.tight_layout()
    fig.savefig(QA / "k_curve.png")
    plt.close(fig)
    m = read(OUT / "fold_verification.json")["metrics"]
    fig, ax = plt.subplots(figsize=(6.5, 1.6))
    vals = [0.4*m["ic_mean"], 0.3*m["annual_excess"], 0.3*(1-m["mean_turnover"])]
    left = 0
    for value, label, color in zip(vals, ["排序相关项", "超额收益项", "名单稳定项"], ["#607e96", "#a36c4b", "#687c68"]):
        ax.barh([0], [value], left=left, color=color, height=0.42, label=label)
        ax.text(left+value/2, 0, f"{value:.4f}", ha="center", va="center", color="white", fontsize=10)
        left += value
    ax.set(xlim=(0, 0.46), xlabel="官方综合分的加权贡献", yticks=[])
    handles, labels = ax.get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", bbox_to_anchor=(0.51, 0.99), ncol=3, frameon=False)
    fig.subplots_adjust(left=0.05, right=0.99, bottom=0.32, top=0.70)
    fig.savefig(QA / "score_terms.png")
    plt.close(fig)
    fig, ax = plt.subplots(figsize=(6.5, 2.1))
    ax.axis("off")
    labels = ["量价历史\n63 / 67 特征", "Top-K 专门目标\n候选成员来源", "XGBoost 与 W5a\n共识与排序", "留存约束\n日度分位", "训练期映射\n下一日收益预测"]
    for i, label in enumerate(labels):
        xpos = i * 0.2 + 0.095
        ax.text(xpos, 0.54, label, ha="center", va="center", transform=ax.transAxes,
                bbox={"facecolor": "white", "edgecolor": "#777777", "boxstyle": "round,pad=0.55"}, fontsize=8.7)
        if i < 4:
            ax.annotate("", xy=(xpos+0.112, 0.54), xytext=(xpos+0.083, 0.54), xycoords="axes fraction",
                        arrowprops={"arrowstyle": "->", "color": "#666666"})
    ax.text(0.5, 0.12, "模型与参数在训练期选定；测试日仅使用当日及此前量价；收益标签只用于事后披露", ha="center", transform=ax.transAxes, fontsize=8)
    fig.tight_layout()
    fig.savefig(QA / "flow.png")
    plt.close(fig)


def build():
    selection = read(OUT / "selection.json")
    verify = read(OUT / "fold_verification.json")
    retro = read(OUT / "retrospective.json")["metrics"]
    checks = read(OUT / "submission_checks.json")
    assert verify["status"] == checks["status"] == "passed"
    charts()
    doc = Document(REF)
    # This is a full content rewrite. Retain the reference's styles, margins and page furniture.
    for element in list(doc._element.body):
        if element.tag != qn("w:sectPr"):
            doc._element.body.remove(element)
    for section in doc.sections:
        for region in (section.header, section.footer):
            for p in region.paragraphs:
                p.clear()
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Caption"):
        style = doc.styles[name]
        style.font.name = "宋体"
        style.font.size = Pt(10.5)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
    for node in doc.styles.element.findall(".//" + qn("w:pBdr")):
        node.getparent().remove(node)

    def para(text, style="Normal", indent=True, bold=False):
        p = doc.add_paragraph(style=style)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.widow_control = True
        snap = OxmlElement("w:snapToGrid")
        snap.set(qn("w:val"), "0")
        p._p.get_or_add_pPr().append(snap)
        if style == "Normal" and indent:
            p.paragraph_format.first_line_indent = Pt(21)
            p._p.get_or_add_pPr().find(qn("w:ind")).set(qn("w:firstLineChars"), "200")
        r = p.add_run(text)
        r.font.name = "宋体"
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.bold = bold or style.startswith("Heading")
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
        r._element.rPr.rFonts.set(qn("w:cs"), "宋体")
        return p

    def heading(text):
        return para(text, "Heading 1", False)

    def subheading(text):
        return para(text, "Heading 2", False)

    def next_page(title):
        p = heading(title)
        p.paragraph_format.page_break_before = True

    def table(caption, headers, rows, widths):
        para(caption, "Caption", False)
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for i, text in enumerate(headers):
            t.rows[0].cells[i].text = text
        repeat = OxmlElement("w:tblHeader")
        t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells = t.add_row().cells
            for i, text in enumerate(row):
                cells[i].text = str(text)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement("w:" + side)
            b.set(qn("w:val"), "single")
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "D9D9D9")
            borders.append(b)
        t._tbl.tblPr.append(borders)
        for ir, row in enumerate(t.rows):
            for ic, cell in enumerate(row.cells):
                cell.width = Cm(widths[ic])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side in ("top", "bottom", "left", "right"):
                    el = OxmlElement("w:"+side)
                    el.set(qn("w:w"), "60" if side in ("top", "bottom") else "90")
                    el.set(qn("w:type"), "dxa")
                    margins.append(el)
                tcpr.append(margins)
                if ir == 0:
                    fill = OxmlElement("w:shd")
                    fill.set(qn("w:fill"), "E7E7E7")
                    tcpr.append(fill)
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.5
                    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.first_line_indent = Pt(0)
                    p.paragraph_format.left_indent = p.paragraph_format.right_indent = Pt(0)
                    p._p.get_or_add_pPr().find(qn("w:ind")).set(qn("w:firstLineChars"), "0")
                    snap = OxmlElement("w:snapToGrid")
                    snap.set(qn("w:val"), "0")
                    p._p.get_or_add_pPr().append(snap)
                    if ic > 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "宋体"
                        r.font.size = Pt(10.5)
                        r.bold = ir == 0
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
                        r._element.rPr.rFonts.set(qn("w:cs"), "宋体")
        return t

    def figure(name, caption):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(QA / name), width=Cm(16.5))
        para(caption, "Caption", False)

    p = para("基于量价信号与分工集成的\n股票收益预测", "Title", False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(150)
    p.runs[0].font.size = Pt(20)
    p.runs[0].bold = True
    p = para("连接验证选模与收益预测数据分析报告", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    p = para("小组名称  因子标尺小组", indent=False)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(78)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    body = doc.sections[-1]
    body.footer.is_linked_to_previous = False
    p = body.footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    page_num = body._sectPr.find(qn("w:pgNumType"))
    if page_num is None:
        page_num = OxmlElement("w:pgNumType")
        body._sectPr.append(page_num)
    page_num.set(qn("w:start"), "1")

    # Body page 1
    heading("摘要")
    para("赛题要求根据原始日频量价数据预测下一交易日股票收益。本研究在2018至2024年训练数据内，按2020至2024年逐年向前验证，以连接样本外预测的官方综合分选择完整方案。选中模型由Top-K成员信号、树模型与GRU集成排序、名单留存约束和单调收益率映射组成。模型分别服务于头部收益、全截面排序和名单连续性，再输出具有训练期收益刻度的预测值。")
    para("选中方案的连接验证综合分为0.445916，Rank IC为0.118068，Top10%算术年化超额收益为36.10%，平均Jaccard换手为3.20%。相对原67特征16种子XGBoost完整方案，综合分提高0.018508。Top-K训练损失的最优比例为17.5%，而官方评分与最终名单仍使用前10%，两者承担不同作用。")
    para(f"冻结方案后，对写出的提交CSV开展测试期回顾评分，综合分为{retro['final_score']:.6f}。测试结果只作事后披露。历史验证曾被反复查看，选中分数存在累积选择偏差，不能视为未来表现保证。提交件覆盖1,599,600条记录，收益率映射在训练期拟合，写入后仍保持每日排序。")
    para("关键词：量价信号；Top-K目标；分工集成；向前验证；收益率映射", indent=False)
    heading("1 背景介绍与研究目标")
    para("股票下一日收益噪声较大，排序、头部组合表现和候选名单稳定性并不总能同时改善。赛题以三个分项共同计分，因此模型需要评价最终输出方案，而非仅比较原始回归误差或某一个分项。研究目标是利用允许的量价字段，在训练数据内选出官方综合分最高且能生成完整预测的方案[1]。")
    table("表1 官方评分目标及样本范围", ["分项", "权重", "官方计算范围"], [
        ["平均Rank IC", "40%", "有真实标签的全部股票"],
        ["年化Top10%超额", "30%", "非涨停且有标签的股票"],
        ["1减Jaccard换手", "30%", "非涨停股票 不按标签筛选"]], [4.1, 1.9, 10.5])
    para("综合分等于0.4×日均Rank IC＋0.3×日均头部超额×252＋0.3×(1－日均换手)。交易成本、回撤、显著性和最差年度均可作诊断，但不改变官方权重，也不额外作为本次选模淘汰条件。")

    # Body page 2
    next_page("2 数据介绍与描述分析")
    table("表2 数据规模与缺失覆盖", ["项目", "训练集", "测试输入"], [
        ["日期", "20180102至20241231", "20250102至20260608"],
        ["股票数", "4,650", "4,650"], ["交易日数", "1,699", "344"],
        ["记录数", "7,900,350", "1,599,600"],
        ["有效训练标签", "6,758,531", "题面不提供"],
        ["价格缺失记录", "1,139,980", "44,844"]], [5.0, 5.75, 5.75])
    subheading("2.1 数据字段与缺失处理")
    para("输入为后复权开盘、最高、最低、收盘价，成交量、成交额和涨跌停标记；股票代码与日期构成唯一键。后复权价格可直接计算相对变化。原始质量复核显示，两份数据均无重复键，非缺失价格均为正，成交量与金额未出现负数。训练集零成交量有38,468条，同时标记涨跌停有4条，保留原交易状态，不凭这些字段补造真实收益。")
    para("训练标签缺失1,141,819条，占14.45%。缺失率从2018年的26.23%降至2024年的2.78%，年度覆盖明显变化。赛题将其主要归因于停牌、未上市等，但字段不足以逐条认定原因。训练目标仅使用有标签行；评分时完全依照官方不同分项的缺失处理，不把标签缺失视为已知的未来可交易状态。")
    subheading("2.2 收益分布与建模含义")
    para("全量有效训练标签的日收益均值为0.0325%，标准差2.9870%，中位数为0；1%与99%分位分别为－8.3482%和10.0000%。正、负收益样本占47.41%和49.02%。单日波动远大于均值，单纯押注上涨不能替代同日相对选股。因此全截面排序臂采用收益秩，而成员来源模型直接学习头部收益相关目标。")
    para("这些统计量来自全部有效原始训练标签，未先截尾。辅助模型中的截尾属于学习目标构造，与描述统计和官方评分的真实标签处理相互独立。树模型可处理历史窗口不足的特征空值；神经网络采用训练窗口内确定的预处理，不以未来行情向前填充。")
    subheading("2.3 信号形成时点")
    para("特征可使用预测日收盘及更早行情，故信号解释为盘后预测下一交易日。最终拟合包括2024年12月31日标签，该标签到2025年1月2日收盘才成熟；模型最早在首个测试日盘后使用，不能被描述为2024年底即可完成全部训练。比赛收盘到收盘评分与真实下单可获得价格应分别理解。")

    # Body page 3
    next_page("3 模型输入与集成结构")
    subheading("3.1 基础量价与同伴特征")
    para("基础输入共63项，取自赛题允许字段。价格和成交量主要转化为比值、变化率、波动与同日相对位置，窗口涵盖5、10、20、60个交易日。历史运算按股票分组并沿时间方向计算；同日截面统计可在收盘后形成。特征先后顺序与保存模型保持一致。")
    table("表3 输入特征的组成", ["特征类别", "项数", "代表信息"], [
        ["K线与相对价格", "12", "实体 振幅 上下影线 相对收盘价"],
        ["滚动量价", "20", "价格变化 均线 价格与成交量波动"],
        ["动量和流动性", "25", "多尺度收益 量额比 高点位置 交易状态"],
        ["同日截面百分位", "6", "收益 振幅 成交额 波动的相对名次"],
        ["同伴信息", "4", "同组收益 滞后值 三日均值 自身偏离"]], [4.5, 1.6, 10.4])
    para("67项输入在基础特征之上增加4项同伴信号。每日按20日收益和相对价格波动分别划分五组，交叉形成至多25个量价单元，排除自身后计算同组其他股票的当日平均收益，再生成一日滞后、三日均值和自身偏离。同伴代表量价状态相近，并非行业关系或因果传导。")
    subheading("3.2 集成成员与角色划分")
    para("原67特征XGBoost集成使用16个种子，各模型预测先转为同日百分位再等权平均。W5a由原LightGBM、单个XGBoost、两日目标树模型、剪枝特征树模型及20日窗口64隐藏单元GRU组成，五路预测等权秩平均。最终排序臂再将W5a与16种子模型按0.75比0.25融合。")
    figure("flow.png", "图1 从历史量价到最终收益预测的处理流程")
    para("不同组件保留不同任务：Top-K目标提供成员来源，16种子模型在提升块内形成共识，W5a与树集成提供当日相对次序。收益辅助模型仅调整特定头部位置。比较以完整输出流程为单位，避免拿神经网络原始预测与优化后的树模型输出直接判断优劣。")

    # Body page 4
    next_page("4 头部成员目标与训练方案")
    subheading("4.1 直接面向头部收益的Top K损失")
    para("成员来源模型采用LightGBM自定义成对logistic目标。每个训练日按当前预测选择前K个样本，从其余样本中为每个头部成员随机抽取8个对比对象，以真实收益差的绝对值加权成对损失。抽样按种子与轮数确定，日组归一化在训练开始前固定，目标保留原始Float64收益标签。")
    para("训练K等于当日有效训练行数乘比例后取整。比例17.5%由训练期完整连接验证选中；官方收益和换手仍取前10%。训练梯度覆盖与评分的名单规模并不要求相同。对头部与其他样本换名的收益差提供了学习依据，但代理损失并不等于官方总分，最终仍须经过完整输出验证。")
    table("表4 主要训练与后处理参数", ["模块", "固定设置"], [
        ["Top-K成员来源", "63特征 180轮 学习率0.05 31叶 seed2026 K比例0.175"],
        ["Top-K正则与采样", "最小叶300 行列采样0.85 L2为5 max_bin127"],
        ["67特征XGBoost", "16种子 360轮 学习率0.025 深度7"],
        ["成员来源后处理", "平滑0.9 持续奖励0.3 头部辅助权重0.25"],
        ["排序臂后处理", "平滑0 持续奖励0 头部辅助权重0.25"],
        ["最终留存", "昨日名单至少保留round(0.985×当日头部人数)"]], [4.1, 12.4])
    subheading("4.2 不同目标族的比较")
    para("近期实验同时检验LambdaRank、rank_xendcg和83项扩展输入。容量匹配的LambdaRank graded且截断500方案得到0.437045，83特征Top-K得到0.443188，均低于选中方案。排序目标的相关度使用整数等级；早期浮点相关度导致两个rank_xendcg臂结果相同的运行已作废，修订后的结果单独保留。")
    para("部分目标提高了成员来源模型的全局IC，却降低头部超额。该现象说明在本次已测候选中，成员来源与全截面排序的优化任务存在差异；它是有限实验的观察，不能推导成所有模型都满足的规律。扩展特征也未自动带来更高完整得分。")
    subheading("4.3 训练样本与信息边界")
    para("Top-K训练剔除无效价格、涨停、原始标签非有限的行，并移除不足100只股票的训练日。逐年验证仅使用验证年前的数据，排除跨边界尚未成熟的一日标签。参数、种子、特征、集成、后处理和校准方案均在训练期确定；最终全训练期拟合不读取测试标签。")

    # Body page 5
    next_page("5 最终名单与收益率输出")
    subheading("5.1 共识成员来源")
    para("首先在非涨停股票中按Top-K成品信号取每日前十分位，作为提升块；将该块置于其他股票之上，块内由67特征16种子成品信号排序。由此获得一个兼顾收益目标和树集成共识的成员分数。提升只读取当日预测与涨停标记，不读取未来收益、是否缺失标签或后续日期。")
    subheading("5.2 精确基数与名单留存")
    para("在成员分数之和最大且人数固定的条件下，最终头部名单至少保留昨日实际头部中的round(0.985×当日头部人数)只。如果前日成员今日可参与股票不足，则下调到可参与人数并由成员分数补齐。首日冷启动；年度连接处继续使用上一日名单，评分保留全部四个跨年换手观测。")
    para("选中名单取得当日可参与股票最大的分数槽位，内部按最终排序臂排列；其余股票按同一排序臂排列。涨停股保留原排序分数，因为它们虽不进入收益和换手名单，仍可进入Rank IC。股票代码确定的微小偏移用于处理并列，未使用标签打破并列。")
    subheading("5.3 两种平滑服务于不同角色")
    para("成员来源保留0.9的历史平滑及0.3的持续性奖励，用于确定持续候选；排序臂使用平滑0、持续奖励0，保留当期信息。辅助收益目标在训练期逐日1%与99%截尾、中心化，再按训练折整体均方根缩放；融合权重为0.25。排序臂自身换手无需独立优化，最终名单换手由留存构造控制。")
    subheading("5.4 从排序转换为收益预测")
    para("最终预测先转为同日平均秩分位，在2020至2024年样本外预测与已实现收益上，以2,000个分位区间拟合单调isotonic映射。输出为映射值加10的负9次方乘分位，用微小严格递增项保持排序和原有并列类。该映射给出了训练期条件收益刻度，pred可以解释为下一日收益预测，但未经独立校准评估，不能将幅度视为可靠保证。")
    para("最终映射在2025年预测前使用全部训练期OOF标签拟合；历史映射后评分仅证明与原排序分数完全保序，不能被描述为各历史年度当时已经拥有这一收益刻度。官方三项指标只依赖同日次序，已实测映射前后综合分均为0.445916。")
    para("CSV写出精度设为17位。单调映射可能含平坦段，微小分位项决定相邻名次；精度不足会把不同预测舍入为并列，改变头部边界。提交检查因此读取实际CSV，再核对每日秩、键顺序、重复键和有限值，而非只检查内存数组。")

    # Body page 6
    next_page("6 连接验证与模型选择")
    para("验证为2020至2024年五个逐年扩展训练窗口，连接覆盖1,212个交易日、5,635,800行。每年模型只拟合此前可获得的数据；原始模型后处理按冻结年度协议冷启动，最终留存按连续日期运行。官方评分器一次计算连接区间，五年等权年度均分只保留作诊断。")
    table("表5 完整方案的同口径连接综合分", ["候选", "综合分", "与选中差"], [
        ["原67特征16种子完整方案", "0.427408", "－0.018508"],
        ["共识与融合 Top-K10%", "0.443191", "－0.002726"],
        ["共识与融合 Top-K20%", "0.445092", "－0.000824"],
        ["共识与融合 Top-K17.5%", "0.445916", "0"],
        ["83特征Top-K10%", "0.443188", "－0.002728"]], [10.1, 3.2, 3.2])
    figure("k_curve.png", "图2 不同训练Top-K比例的完整连接验证得分")
    para("汇总744条已评分记录，包含重复候选和对照，按训练期官方连接分降序，17.5%为最高。原结论曾以局部曲线波动和拆分半迁移为由保留10%；本次按既定标准修正选择，额外研究诊断不再替代最高分规则，也不将需重训当作不可交付。")
    figure("score_terms.png", "图3 选中方案的三项官方加权贡献")
    para("选中方案前、后段得分为0.448512和0.442018；相对10%方案前段提高0.005807，后段下降0.001909。增益具有时期差异，反复搜索的冠军分数不能视为独立泛化证据。")

    # Body page 7
    next_page("7 应用验证与预测交付")
    subheading("7.1 训练分数与预测流程复核")
    para("保存的五个Top-K17.5%原始验证预测逐折核对股票日期键，与训练边界元数据一致。完整重构得到0.4459160901521865，与冻结记录完全一致。收益映射后的官方评分亦一致。原63特征Top-K训练路径在2022折曾逐位复现冻结预测，作为训练代码的参照证明。")
    para("最终成员来源用允许的2018至2024年标签训练180轮并保存模型，其他组件使用已核验权重或已证明可重构的预测。W5a的GRU原始列已重建，并在原后处理下逐元素复现历史成品。每个中间数组保留键对应关系与SHA256，最终提交件另设写后检查。")
    table("表6 提交件的实际检查结果", ["检查", "结果"], [
        ["字段", "ts_code trade_date pred"], ["记录数", "1,599,600"],
        ["覆盖", "344交易日 4,650股票"], ["日期范围", "20250102至20260608"],
        ["空值与重复键", "均为0"], ["有效数值与原输入顺序", "全部通过"],
        ["收益映射与CSV保序", "每日秩完全一致"]], [7.0, 9.5])
    subheading("7.2 测试区间的事后披露")
    para(f"模型和提交文件冻结后，未经修改的官方evaluate.py直接读取写出的CSV，在现有回顾标签区间得到综合分{retro['final_score']:.6f}，Rank IC为{retro['ic_mean']:.6f}，年化超额为{retro['annual_excess']*100:.2f}%，平均Jaccard换手为{retro['mean_turnover']*100:.2f}%。这些测试标签已被历史研究查看，结果属于回顾比较，未用于改变本次选择。")
    historical_best = 0.417099224
    if retro["final_score"] > historical_best:
        para("在已有实际CSV官方回顾评分记录中，本次冻结提交件得分最高；这里报告的是已评分记录范围，不推断全部未评分候选的测试表现。验证选中方案与该记录范围内测试最高方案在本次一致。")
    else:
        para("在已有实际CSV官方回顾评分记录中，波13的W5a单排序臂方案得分最高，为0.417099，原67特征16种子方案为0.386058。验证选中方案与已知测试最高方案分别报告；测试分数不回流选模，未评分候选也不能被认定优劣。")
    para("预测文件覆盖全部344日，现有官方回顾标签覆盖至2026年6月5日的343日，两者范围不同。隐藏或未成熟标签不补造。保序校准改善输出的语义，不会独立提高官方排名分数。真实使用还须评价收益刻度、可执行成交价、费用、冲击和容量。")

    # Body page 8
    next_page("8 产品思路与总结结论")
    subheading("8.1 盘后候选研究工作台")
    para("产品面向需要重复筛选全市场股票的投研团队。每日收盘后接收量价字段，先检查缺价、零成交与异常标记，再输出下一日收益预测、截面名次、头部名单及进入退出记录。研究员可查看个股与量价相似股票的差异，区分预测名次变化和名单留存约束造成的变化。")
    table("表7 产品模块及验证重点", ["模块", "使用信息", "验证重点"], [
        ["数据与信号", "质量标记 收益预测 排名", "时点与前缀一致性"],
        ["候选名单", "成员来源 名次变化 留存原因", "规则解释与模型一致"],
        ["效果监控", "排序表现 幅度误差 模拟成交", "费用 滑点 无法成交"],
        ["版本管理", "冻结配置 文件哈希 回退记录", "复现与人工复核"]], [3.2, 7.1, 6.2])
    para("部署前先进行影子运行，固定方案后记录新时期预测，采用下一日实际可获得的成交价格，保留停牌、涨跌停和无法成交记录。官方Jaccard名单换手不等于资金换手，也不能直接代表交易费用。收益率映射虽具有训练期条件收益含义，仍需用后续标签评估幅度偏差；界面应同时呈现相对名次与预测幅度。")
    para("商业形式可以是研究终端订阅或机构内部数据接口，价值在于统一筛选口径、减少重复工作、提供可追溯记录。目前没有客户转化或经营收益数据，商业效果需要试点验证。模型更新先在训练或冻结验证期内比较，替换时保留前一版本与完整哈希清单。")
    subheading("8.2 研究结论")
    para("本研究将下一日收益预测拆为成员来源、全截面排序、名单连续性和收益刻度四个相互连接的部分。Top-K目标负责候选来源，量价树模型与GRU秩集成负责当日排序，精确留存规则控制名单变化，训练期单调映射形成收益预测。该结构在统一连接验证下得到0.445916，比原方案提高0.018508。")
    para("近期LambdaRank、rank_xendcg、83特征扩展及角色替换未超过17.5%Top-K方案。旧结论中的额外研究门槛与项目选模规则不一致，已依最高官方连接分完成修正、最终拟合、提交检查与报告更新。改善不是增加种子或特征的单调结果，应评价各组件在完整输出中的作用。")
    para("已观察时期的搜索收益不能保证未来表现。反复选模、后段相对退化、毛收益评分与收益刻度迁移都是实际限制。下一阶段应固定本方案，在新时期检查预测方向、收益幅度与可执行净收益，决定它是否适合投研服务或交易试点。")

    # References page, outside the eight-page body.
    next_page("参考文献")
    for text in [
        "[1] 主办方 赛题五 基于原始量价数据的股票收益预测 更新版及官方Python评价脚本evaluate.py 本地赛题资料。",
        "[2] 项目 后续选模标准 训练内选模 冻结完整方案后预测 以连接样本外官方综合分为主分数。",
        "[3] 项目 连接验证实验 第1至39波冻结协议和结果 2026年10月2日至3日。",
    ]:
        para(text, indent=False)
    heading("附录 数据与复核依据")
    para("数据规模、缺失率和收益统计承接原报告对原始赛题CSV的全量复核。表5及图2直接读取冻结连接验证结果，图3读取本次复现的官方指标。模型结构由对应训练、后处理及构造源码核对。报告不把本地历史年度等权均分替代为连接评分，也不将回顾结果标成独立盲测。")
    para("本次选择覆盖744条已评分记录，其中有重复候选及对照，因此不等于744个独立模型。选择冻结记录为selection.json，训练边界与重构记录为fold_verification.json，最终模型记录为train_test.json，提交实检为submission_checks.json，测试回顾披露为retrospective.json。")
    para("运行顺序为select、folds、train、build、retrospective。最后一步单独读取回顾标签，仅用于报告。提交文件写后秩与内存秩一致，三个评分分项的分母分别计算，首日换手无观测，年度间换手不断开。")
    para("最终交付文件为submission.csv和本报告的PDF及可编辑Word版本。提交平台的实际上传属于比赛操作，不由本地交付构建触发。文件哈希和版本指针随交付清单保存。")
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = "基于量价信号与分工集成的股票收益预测"
    doc.save(TARGET)
    artifact = {"reference": str(REF), "reference_sha256": hashlib.sha256(REF.read_bytes()).hexdigest(),
                "output": str(TARGET), "page_system": "A4 portrait, top1.9 bottom1.8 left/right2.15cm",
                "body": "8 planned pages, Song 10.5pt, 1.5 spacing, first indent2 characters",
                "scope": "Full content rewrite; retain academic styles and page geometry. Cover plus eight body pages plus references.",
                "tables": len(doc.tables), "figures": len(doc.inline_shapes)}
    (QA / "artifact.md").write_text("# Report contract\n\n"+json.dumps(artifact, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    print(TARGET)


if __name__ == "__main__":
    build()
